from playwright.sync_api import sync_playwright
import openpyxl
from pathlib import Path
import re
import time
import argparse

# ===== DEFAULT CONFIG =====
ROOT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_FILE = ROOT_DIR / "Assignment 1 - Test cases.xlsx"
DEFAULT_URL = "https://www.pixelssuite.com/chat-translator"

# ===== REQUIRED HEADERS =====
REQUIRED_COLUMNS = [
    "TC ID",
    "Input length type",
    "Input",
    "Expected output",
    "Actual Output",
    "Status",
    "Singlish input types covered",
    "Evidence or rationale for the input type covered"
]


def normalize(text):
    if not text:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(text).lower())


def ensure_columns(ws):
    headers = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]
    col_map = {}
    for col_name in REQUIRED_COLUMNS:
        col_index = None
        for i, h in enumerate(headers, 1):
            if h and normalize(col_name) == normalize(h):
                col_index = i
                break
        if not col_index:
            col_index = ws.max_column + 1
            ws.cell(1, col_index).value = col_name
        col_map[col_name] = col_index
    return col_map


def get_output(page):
    try:
        boxes = page.locator("textarea")
        if boxes.count() >= 2:
            val = boxes.nth(1).input_value().strip()
            if val and "Disclaimer" not in val:
                return val
    except:
        pass
    return ""


def normalize_sinhala(text):
    """Normalize Sinhala text for comparison — strip whitespace and collapse spaces."""
    if not text:
        return ""
    return re.sub(r"\s+", " ", str(text).strip())


def is_pass(expected, actual):
    """Compare expected vs actual output — pass if they match after normalization."""
    return normalize_sinhala(expected) == normalize_sinhala(actual)


def main():
    parser = argparse.ArgumentParser(description="Playwright test automation for Singlish→Sinhala translator")
    parser.add_argument("--excel", default=str(DEFAULT_INPUT_FILE), help="Path to the Excel test cases file")
    parser.add_argument("--url", default=DEFAULT_URL, help="URL of the translator app")
    parser.add_argument("--wait-ms", type=int, default=5000, help="Wait time (ms) for page after input (default: 5000)")
    parser.add_argument("--type-delay-ms", type=int, default=80, help="Delay (ms) between keystrokes (default: 80)")
    parser.add_argument("--slow-mo-ms", type=int, default=200, help="Playwright slow-mo delay in ms (default: 200)")
    parser.add_argument("--save-every", type=int, default=1, help="Save Excel every N rows (default: 1)")
    parser.add_argument("--keep-open", action="store_true", help="Keep browser open after tests finish")
    parser.add_argument("--max-rows", type=int, default=50, help="Maximum number of test rows to run (default: 50)")
    parser.add_argument("--headless", action="store_true", help="Run browser in headless mode")
    args = parser.parse_args()

    input_file = Path(args.excel)
    output_file = input_file

    print(f" Excel file : {input_file}")
    print(f" URL        : {args.url}")
    print(f"  Wait ms    : {args.wait_ms}")
    print(f" Type delay : {args.type_delay_ms}ms")
    print(f" Slow-mo    : {args.slow_mo_ms}ms")
    print()

    wb = openpyxl.load_workbook(input_file)
    ws = wb.active
    col = ensure_columns(ws)

    input_col    = col["Input"]
    expected_col = col["Expected output"]
    actual_col   = col["Actual Output"]
    status_col   = col["Status"]

    max_row = min(ws.max_row, args.max_rows + 1)  # +1 for header row

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=args.headless,
            slow_mo=args.slow_mo_ms
        )
        page = browser.new_page()
        page.goto(args.url)
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(2000)

        input_box = page.locator("textarea, input").first
        pass_count = 0
        fail_count = 0
        error_count = 0

        for row in range(2, max_row + 1):
            text     = ws.cell(row, input_col).value
            expected = ws.cell(row, expected_col).value

            if not text:
                continue

            tc_id = ws.cell(row, 1).value or f"Row {row}"
            print(f"➡  [{tc_id}] Row {row} | Input: {str(text)[:60]}")

            try:
                # Clear and type input
                input_box.click()
                input_box.fill("")
                input_box.type(str(text), delay=args.type_delay_ms)
                page.keyboard.press("Enter")

                # Wait for output to appear
                actual = ""
                checks = args.wait_ms // 300
                for _ in range(checks):
                    actual = get_output(page)
                    if actual:
                        break
                    time.sleep(0.3)

                if not actual:
                    actual = "NO OUTPUT"

                ws.cell(row, actual_col).value = actual

                # Determine pass/fail
                if actual == "NO OUTPUT":
                    status = "FAIL"
                    fail_count += 1
                elif is_pass(expected, actual):
                    status = "PASS"
                    pass_count += 1
                else:
                    status = "FAIL"
                    fail_count += 1

                ws.cell(row, status_col).value = status
                print(f"   {'✅' if status == 'PASS' else '❌'} {status} | Actual: {actual[:60]}")

            except Exception as e:
                ws.cell(row, actual_col).value = "ERROR"
                ws.cell(row, status_col).value = "FAIL"
                fail_count += 1
                error_count += 1
                print(f"   ❌ ERROR: {e}")

            # Save periodically
            if (row - 1) % args.save_every == 0:
                wb.save(output_file)

        # Final save
        wb.save(output_file)

        print()
        print("=" * 60)
        print(f"✅ PASS  : {pass_count}")
        print(f"❌ FAIL  : {fail_count}")
        print(f"ERRORS: {error_count}")
        print(f"TOTAL : {pass_count + fail_count}")
        print("=" * 60)
        print(f"Results saved to: {output_file}")

        if args.keep_open:
            print("🔓 Browser kept open. Close it manually when done.")
            input("Press Enter to close browser and exit...")

        browser.close()

    print("DONE")


if __name__ == "__main__":
    main()